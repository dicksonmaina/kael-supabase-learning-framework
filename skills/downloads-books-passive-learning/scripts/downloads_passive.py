#!/usr/bin/env python3
"""Downloads Books Passive Learning."""

from __future__ import annotations

import argparse
import json
import random
import re
import subprocess
import sys
from pathlib import Path

CORPUS = Path(__file__).resolve().parent.parent.parent.parent / "study-materials" / "downloads-books" / "corpus"
BUILD = Path(__file__).resolve().parent.parent.parent.parent / "study-materials" / "downloads-books" / "build"
UNIFIED = BUILD / "unified.md"
MANIFEST = BUILD / "manifest.json"
SOURCES = BUILD / "sources.json"


def ensure_dirs() -> None:
    CORPUS.mkdir(parents=True, exist_ok=True)
    BUILD.mkdir(parents=True, exist_ok=True)


def iter_corpus_files():
    if not CORPUS.exists():
        return []
    return sorted([p for p in CORPUS.iterdir() if p.is_file()])


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            import tempfile

            with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tf:
                tf.close()
                res = subprocess.run(
                    ["pdftotext", "-layout", str(path), tf.name],
                    capture_output=True,
                    text=True,
                )
                if res.returncode == 0:
                    text = Path(tf.name).read_text(errors="ignore")
                else:
                    text = ""
                Path(tf.name).unlink(missing_ok=True)
                return text
        elif suffix == ".docx":
            from docx import Document

            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        elif suffix == ".pptx":
            try:
                from pptx import Presentation
            except ImportError:
                return "[pptx extraction unavailable: install python-pptx]"
            prs = Presentation(str(path))
            parts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        parts.append(shape.text.strip())
            return "\n\n".join(parts)
        else:
            return path.read_text(errors="ignore")
    except Exception as exc:
        return f"[ERROR extracting {path.name}: {exc}]"


def cmd_ingest(_: argparse.Namespace) -> int:
    ensure_dirs()
    files = iter_corpus_files()
    if not files:
        print("No corpus files found.")
        return 0

    manifest = {"files": [], "total_files": len(files)}
    sources = {"sources": []}
    chunks = []
    for path in files:
        text = extract_text(path)
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if not text:
            text = f"[No extractable text from {path.name}]"
        chunks.append(f"# {path.name}\n\n{text}\n")
        manifest["files"].append(
            {
                "name": path.name,
                "size_bytes": path.stat().st_size,
                "type": path.suffix.lower().lstrip("."),
            }
        )
        sources["sources"].append(
            {"file": path.name, "path": str(path), "size_bytes": path.stat().st_size}
        )

    UNIFIED.write_text("\n---\n".join(chunks), encoding="utf-8")
    MANIFEST.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    SOURCES.write_text(json.dumps(sources, indent=2), encoding="utf-8")
    print(f"Ingested {len(files)} files -> {UNIFIED}")
    return 0


def cmd_stats(_: argparse.Namespace) -> int:
    files = iter_corpus_files()
    by_type: dict[str, int] = {}
    for p in files:
        by_type[p.suffix.lower()] = by_type.get(p.suffix.lower(), 0) + 1
    print(f"Total files: {len(files)}")
    for k, v in sorted(by_type.items()):
        print(f"  {k}: {v}")
    if MANIFEST.exists():
        print(f"Manifest: {MANIFEST}")
    if UNIFIED.exists():
        print(f"Unified size: {UNIFIED.stat().st_size / 1024:.1f} KB")
    return 0


def cmd_review(_: argparse.Namespace) -> int:
    files = iter_corpus_files()
    if not files:
        print("No corpus files found.")
        return 0
    choice = random.choice(files)
    print(f"# {choice.name}\n")
    text = extract_text(choice)
    print(text[:4000])
    return 0


def cmd_search(args: argparse.Namespace) -> int:
    if not UNIFIED.exists():
        print("Missing unified.md; run ingest first.")
        return 1
    text = UNIFIED.read_text(encoding="utf-8", errors="ignore")
    pattern = re.compile(re.escape(args.query), re.IGNORECASE)
    matches = pattern.findall(text)
    print(f"Found {len(matches)} matches for '{args.query}'.")
    return 0


def cmd_cards(args: argparse.Namespace) -> int:
    files = iter_corpus_files()
    if not files:
        print("No corpus files found.")
        return 0
    count = min(args.count, len(files))
    print(f"Showing {count} sample cards from corpus:")
    for path in random.sample(files, count):
        text = extract_text(path)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        print(f"\n## {path.name}")
        for line in lines[:5]:
            print(f"- {line}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Downloads Books Passive Learning")
    parser.add_argument("command", choices=["ingest", "review", "search", "cards", "stats"])
    parser.add_argument("query", nargs="?", help="Search query for search command")
    parser.add_argument("--count", type=int, default=10, help="Card count for cards command")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    if args.command == "search" and not args.query:
        parser.error("search requires a query")
    handlers = {
        "ingest": cmd_ingest,
        "review": cmd_review,
        "search": cmd_search,
        "cards": cmd_cards,
        "stats": cmd_stats,
    }
    return handlers[args.command](args)


if __name__ == "__main__":
    sys.exit(main())
